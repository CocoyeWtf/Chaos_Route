"""Ticket #78 : périmètre de l'export « Tournées ERT ».

L'export doit sortir EXACTEMENT les tournées de la vue d'où il part :
- `source=ordonnancement` → date de PLANIFICATION (`Tour.date`), et rien d'autre :
  c'est ce que charge l'onglet ordonnancement (`/tours/?date=`, sans filtre de
  statut ni d'heure). Les brouillons sans contrat ni heure en font donc partie —
  le service transport veut un export « brut » des tournées construites (3e passe
  du ticket). Seule reste exclue la tournée d'un AUTRE jour de planification.
- `source=postier` (défaut) → date de LIVRAISON (`Tour.delivery_date`), heure de
  départ posée, statut != DRAFT — comme l'onglet postier.

Historique : l'export filtrait d'abord sur la date de livraison depuis les deux
vues, donc à l'ordonnancement il sortait les tournées de la veille en ratant les
vraies. C'est la date, et non le brouillon, qui était le défaut — d'où le retour
des brouillons dans le périmètre à la demande du demandeur.
"""

import io
import uuid

import pytest
from openpyxl import load_workbook

PLANIF = "2026-09-01"
# Colonne « H.Départ » de la feuille (AT) — importée du module pour ne pas figer
# un indice en dur dans le test. / Departure-time column, taken from the module.
from app.api.exports import _PLANNING_OPS_COL0 as _OPS_DEPARTURE_COL
LIVRAISON = "2026-09-02"


async def _make_base(db_session, region, name="Villers"):
    from app.models.base_logistics import BaseLogistics

    base = BaseLogistics(code=f"B{uuid.uuid4().hex[:5]}", name=name,
                         city=name, region_id=region.id)
    db_session.add(base)
    await db_session.commit()
    await db_session.refresh(base)
    return base


async def _make_tour(db_session, base, *, date, delivery_date, status,
                     departure, priority=None):
    from app.models.tour import Tour, TourStatus

    tour = Tour(
        date=date, delivery_date=delivery_date, code=f"T-{uuid.uuid4().hex[:10]}",
        base_id=base.id, status=getattr(TourStatus, status), priority=priority,
        departure_time=departure, temperature_type="FRAIS",
    )
    db_session.add(tour)
    await db_session.commit()
    await db_session.refresh(tour)
    return tour


def _codes(content: bytes) -> set[str]:
    """Codes tournée (colonne N° Mission) présents dans le classeur."""
    ws = load_workbook(io.BytesIO(content))["Tours"]
    out = set()
    for row in range(7, ws.max_row + 1):
        v = ws.cell(row, 2).value
        if v:
            out.add(str(v))
    return out


@pytest.mark.asyncio
async def test_ordonnancement_exports_planning_date_not_delivery_date(
        client, db_session, test_region):
    """Le coeur du #78 : depuis l'ordonnancement, on exporte la date de planif,
    brouillons compris."""
    base = await _make_base(db_session, test_region)
    # Ordonnancée le 01, livrée le 02 : c'est CE tour que Samuel voit et attend.
    ordo = await _make_tour(db_session, base, date=PLANIF, delivery_date=LIVRAISON,
                            status="DRAFT", departure="06:00", priority=1)
    # Brouillon pur du meme jour, sans heure ni contrat : il fait partie des
    # tournees construites du jour, donc de l'export brut (3e passe).
    brouillon = await _make_tour(db_session, base, date=PLANIF, delivery_date=None,
                                 status="DRAFT", departure=None)
    # Reliquat de la veille, livre le 01 : hors perimetre de la vue ordonnancement.
    veille = await _make_tour(db_session, base, date="2026-08-31",
                              delivery_date=PLANIF, status="VALIDATED",
                              departure="05:00")

    resp = await client.get("/api/exports/postier-planning", params={
        "date": PLANIF, "base_id": base.id, "source": "ordonnancement"})
    assert resp.status_code == 200, resp.text
    codes = _codes(resp.content)
    assert ordo.code in codes
    assert brouillon.code in codes
    assert veille.code not in codes


@pytest.mark.asyncio
async def test_postier_keeps_delivery_date_and_excludes_drafts(
        client, db_session, test_region):
    """Vue postier : date de livraison, et plus de brouillons dans l'export."""
    base = await _make_base(db_session, test_region)
    livre = await _make_tour(db_session, base, date=PLANIF, delivery_date=LIVRAISON,
                             status="VALIDATED", departure="06:00")
    # Meme date de livraison mais encore en brouillon -> l'onglet postier ne
    # l'affiche pas, l'export ne doit plus le sortir non plus.
    draft = await _make_tour(db_session, base, date=PLANIF, delivery_date=LIVRAISON,
                             status="DRAFT", departure="07:00")
    # Livre un autre jour.
    autre = await _make_tour(db_session, base, date=PLANIF, delivery_date=PLANIF,
                             status="VALIDATED", departure="08:00")

    resp = await client.get("/api/exports/postier-planning", params={
        "date": LIVRAISON, "base_id": base.id})
    assert resp.status_code == 200, resp.text
    codes = _codes(resp.content)
    assert livre.code in codes
    assert draft.code not in codes
    assert autre.code not in codes


@pytest.mark.asyncio
async def test_base_id_optional_exports_all_bases(client, db_session, test_region):
    """Le filtre base est masque quand une seule base existe : export sans base."""
    b1 = await _make_base(db_session, test_region, name="Villers")
    b2 = await _make_base(db_session, test_region, name="Gosselies")
    t1 = await _make_tour(db_session, b1, date=PLANIF, delivery_date=LIVRAISON,
                          status="DRAFT", departure="06:00")
    t2 = await _make_tour(db_session, b2, date=PLANIF, delivery_date=LIVRAISON,
                          status="DRAFT", departure="07:00")

    resp = await client.get("/api/exports/postier-planning", params={
        "date": PLANIF, "source": "ordonnancement"})
    assert resp.status_code == 200, resp.text
    codes = _codes(resp.content)
    assert {t1.code, t2.code} <= codes

    # Avec une base : uniquement celle-la.
    resp = await client.get("/api/exports/postier-planning", params={
        "date": PLANIF, "base_id": b1.id, "source": "ordonnancement"})
    codes = _codes(resp.content)
    assert t1.code in codes and t2.code not in codes


@pytest.mark.asyncio
async def test_order_column_blank_when_no_priority(client, db_session, test_region):
    """Colonne « Ordre » = priorite saisie, vide sinon (pas un rang trompeur)."""
    base = await _make_base(db_session, test_region)
    await _make_tour(db_session, base, date=PLANIF, delivery_date=LIVRAISON,
                     status="DRAFT", departure="06:00", priority=4)
    await _make_tour(db_session, base, date=PLANIF, delivery_date=LIVRAISON,
                     status="DRAFT", departure="07:00", priority=None)

    resp = await client.get("/api/exports/postier-planning", params={
        "date": PLANIF, "base_id": base.id, "source": "ordonnancement"})
    ws = load_workbook(io.BytesIO(resp.content))["Tours"]
    assert ws.cell(7, 1).value == 4      # priorite saisie
    assert ws.cell(8, 1).value is None   # non saisie -> vide


@pytest.mark.asyncio
async def test_ordonnancement_exports_raw_drafts_with_empty_ops_columns(
        client, db_session, test_region):
    """Export « brut » (3e passe) : une tournée construite mais pas encore
    ordonnancée sort, avec ses colonnes d'exploitation vides.

    C'est ce vide qui la distingue à la lecture d'une tournée ordonnancée — il
    n'y a donc pas besoin d'une colonne supplémentaire pour le signaler.
    """
    base = await _make_base(db_session, test_region)
    brut = await _make_tour(db_session, base, date=PLANIF, delivery_date=None,
                            status="DRAFT", departure=None)
    ordo = await _make_tour(db_session, base, date=PLANIF, delivery_date=LIVRAISON,
                            status="VALIDATED", departure="06:30")

    resp = await client.get("/api/exports/postier-planning", params={
        "date": PLANIF, "base_id": base.id, "source": "ordonnancement"})
    assert resp.status_code == 200, resp.text
    ws = load_workbook(io.BytesIO(resp.content))["Tours"]

    heures = {}
    for row in range(7, ws.max_row + 1):
        code = ws.cell(row, 2).value
        if code:
            heures[str(code)] = ws.cell(row, _OPS_DEPARTURE_COL).value

    assert brut.code in heures and ordo.code in heures
    assert not heures[brut.code], "la tournée non ordonnancée ne doit pas inventer une heure"
    assert heures[ordo.code] == "06:30"


@pytest.mark.asyncio
async def test_postier_still_excludes_unscheduled(client, db_session, test_region):
    """Garde-fou : l'ouverture aux brouillons ne concerne QUE l'ordonnancement.

    L'onglet postier reste la liste de ce qui part réellement : une tournée sans
    heure de départ n'y a pas sa place.
    """
    base = await _make_base(db_session, test_region)
    sans_heure = await _make_tour(db_session, base, date=PLANIF,
                                  delivery_date=LIVRAISON, status="VALIDATED",
                                  departure=None)
    avec_heure = await _make_tour(db_session, base, date=PLANIF,
                                  delivery_date=LIVRAISON, status="VALIDATED",
                                  departure="06:00")

    resp = await client.get("/api/exports/postier-planning", params={
        "date": LIVRAISON, "base_id": base.id})
    assert resp.status_code == 200, resp.text
    codes = _codes(resp.content)
    assert avec_heure.code in codes
    assert sans_heure.code not in codes


@pytest.mark.asyncio
async def test_remarque_mise_en_evidence(client, db_session, test_region):
    """Ticket #48 : une remarque renseignée ressort en jaune, texte rouge.

    Sur une feuille dense, le postier doit la repérer sans la chercher. Une
    tournée sans remarque garde une cellule neutre — sinon le repère ne servirait
    plus à rien.
    """
    base = await _make_base(db_session, test_region)
    avec = await _make_tour(db_session, base, date=PLANIF, delivery_date=LIVRAISON,
                            status="DRAFT", departure="06:00")
    sans = await _make_tour(db_session, base, date=PLANIF, delivery_date=LIVRAISON,
                            status="DRAFT", departure="07:00")
    avec.remarks = "Attention hayon HS"
    await db_session.commit()

    resp = await client.get("/api/exports/postier-planning", params={
        "date": PLANIF, "base_id": base.id, "source": "ordonnancement"})
    assert resp.status_code == 200, resp.text
    ws = load_workbook(io.BytesIO(resp.content))["Tours"]

    cellules = {}
    for row in range(7, ws.max_row + 1):
        code = ws.cell(row, 2).value
        if code:
            cellules[str(code)] = ws.cell(row, 11)

    marquee = cellules[avec.code]
    assert marquee.value == "Attention hayon HS"
    assert marquee.fill.fgColor.rgb.endswith("FFFF00"), marquee.fill.fgColor.rgb
    assert marquee.font.color.rgb.endswith("FF0000"), marquee.font.color.rgb

    neutre = cellules[sans.code]
    assert not neutre.value
    assert neutre.fill.fgColor.rgb in (None, "00000000"), neutre.fill.fgColor.rgb
