"""Retour des numeros de tournee Infolog dans CMRO (#104) /
Infolog tour numbers fed back into CMRO (#104).

Le numero CMRO n'est pas exploitable par Infolog : c'est Infolog qui attribue
le sien. La boucle ne se refermait pas — CMRO ignorait sous quel numero une
tournee existait en aval, et la tracabilite s'arretait au milieu de la chaine.

On teste le circuit COMPLET (export CMRO -> saisie -> reimport), dans les deux
variantes prevues au ticket : le fichier WMS complete par la macro, et le
fichier ERT rempli a la main par le trafic.
"""

import io
import uuid

import pytest
from openpyxl import load_workbook

HEADER_ROW = 6
FIRST_DATA_ROW = 7


async def _tour(db_session, test_region, test_pdv, *, jour, priorite=1.0):
    from app.models.base_logistics import BaseLogistics
    from app.models.tour import Tour, TourStatus
    from app.models.tour_stop import TourStop

    base = BaseLogistics(code=f"B{uuid.uuid4().hex[:5].upper()}", name="Gosselies",
                         region_id=test_region.id)
    db_session.add(base)
    await db_session.flush()
    tour = Tour(
        date=jour, delivery_date=jour, code=f"T-{uuid.uuid4().hex[:8].upper()}",
        base_id=base.id, status=TourStatus.VALIDATED, driver_name="Jean Dupont",
        departure_time="05:30", priority=priorite,
    )
    db_session.add(tour)
    await db_session.flush()
    db_session.add(TourStop(tour_id=tour.id, pdv_id=test_pdv.id, sequence_order=1, eqp_count=10))
    await db_session.commit()
    await db_session.refresh(tour)
    return tour


def _classeur(contenu: bytes):
    return load_workbook(io.BytesIO(contenu), data_only=True)


def _vers_fichier(wb) -> bytes:
    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()


def _entetes(ws) -> dict:
    return {str(ws.cell(HEADER_ROW, c).value or "").lower(): c
            for c in range(1, ws.max_column + 1)}


def _ligne(ws, entetes: dict, code: str) -> int:
    """Ligne de CETTE tournee dans l'export. / Locate this tour's row."""
    col = entetes["code cmro"]
    for ligne in range(FIRST_DATA_ROW, ws.max_row + 1):
        if str(ws.cell(ligne, col).value or "") == code:
            return ligne
    raise AssertionError(f"tournee {code} absente de l'export")


# ─── Exports : la colonne technique doit etre la ───

@pytest.mark.asyncio
async def test_l_export_wms_porte_le_code_cmro(client, db_session, test_region, test_pdv):
    """Sans identifiant CMRO dans le fichier, le rapprochement au retour
    reposerait sur l'ordre des lignes — c'est-a-dire sur rien."""
    tour = await _tour(db_session, test_region, test_pdv, jour="2026-10-12")
    resp = await client.get("/api/exports/wms-infolog", params={"date": tour.date})
    assert resp.status_code == 200, resp.text
    ws = _classeur(resp.content).active
    codes = {ws.cell(l, 9).value for l in range(1, ws.max_row + 1)}
    assert tour.code in codes
    # Les colonnes que lit la macro ne doivent PAS avoir bouge
    assert ws.cell(1, 2).value is not None       # B = code PDV
    assert ws.cell(1, 7).value is not None       # G = index global


@pytest.mark.asyncio
async def test_l_export_ert_porte_le_code_cmro_et_la_colonne_a_remplir(
    client, db_session, test_region, test_pdv,
):
    tour = await _tour(db_session, test_region, test_pdv, jour="2026-10-13")
    resp = await client.get("/api/exports/postier-planning",
                            params={"date": tour.date, "source": "ordonnancement"})
    assert resp.status_code == 200, resp.text
    ws = _classeur(resp.content)["Tours"]
    entetes = _entetes(ws)
    assert "code cmro" in entetes
    assert "n° infolog" in entetes
    assert ws.cell(_ligne(ws, entetes, tour.code), entetes["code cmro"]).value == tour.code


# ─── Circuit complet ───

@pytest.mark.asyncio
async def test_circuit_ert_le_trafic_saisit_le_numero(client, db_session, test_region, test_pdv):
    """Variante 2 du ticket : le trafic recopie les numeros a la main."""
    tour = await _tour(db_session, test_region, test_pdv, jour="2026-10-14")
    resp = await client.get("/api/exports/postier-planning",
                            params={"date": tour.date, "source": "ordonnancement"})
    wb = _classeur(resp.content)
    ws = wb["Tours"]
    entetes = _entetes(ws)
    ws.cell(_ligne(ws, entetes, tour.code), entetes["n° infolog"], 291186)

    resp = await client.post(
        "/api/imports/tour-numbers",
        files={"file": ("ert.xlsx", _vers_fichier(wb),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    corps = resp.json()
    assert corps["format"] == "ERT"
    assert corps["updated"] == 1
    await db_session.refresh(tour)
    # Excel rend le nombre saisi comme un flottant : il doit ressortir entier
    assert tour.wms_tour_code == "291186"


@pytest.mark.asyncio
async def test_circuit_wms_la_macro_colle_le_numero(client, db_session, test_region, test_pdv):
    """Variante 1 du ticket : la macro Infolog ecrit le numero dans le fichier."""
    tour = await _tour(db_session, test_region, test_pdv, jour="2026-10-15")
    resp = await client.get("/api/exports/wms-infolog", params={"date": tour.date})
    wb = _classeur(resp.content)
    ws = wb.active
    for ligne in range(1, ws.max_row + 1):
        if ws.cell(ligne, 9).value == tour.code:
            ws.cell(ligne, 10, "291187")      # la macro ecrit juste apres

    resp = await client.post(
        "/api/imports/tour-numbers",
        files={"file": ("TMS_vers_wms.xlsx", _vers_fichier(wb),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["format"] == "WMS"
    await db_session.refresh(tour)
    assert tour.wms_tour_code == "291187"


@pytest.mark.asyncio
async def test_le_numero_ecrase_dans_la_colonne_mission_est_compris(
    client, db_session, test_region, test_pdv,
):
    """Habitude du trafic : ecrire le numero dans « N° Mission », la colonne que
    l'oeil cherche. Le code CMRO reste lisible a cote, donc on sait recoller."""
    tour = await _tour(db_session, test_region, test_pdv, jour="2026-10-16")
    resp = await client.get("/api/exports/postier-planning",
                            params={"date": tour.date, "source": "ordonnancement"})
    wb = _classeur(resp.content)
    ws = wb["Tours"]
    entetes = _entetes(ws)
    ws.cell(_ligne(ws, entetes, tour.code), entetes["n° mission"], "291188")

    resp = await client.post(
        "/api/imports/tour-numbers",
        files={"file": ("ert.xlsx", _vers_fichier(wb),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    await db_session.refresh(tour)
    assert tour.wms_tour_code == "291188"


@pytest.mark.asyncio
async def test_le_code_cmro_n_est_jamais_ecrase(client, db_session, test_region, test_pdv):
    """Le code CMRO est la cle technique : feuilles de route, QR d'affectation
    (#98), etiquettes. L'import ne doit toucher QUE le numero Infolog."""
    tour = await _tour(db_session, test_region, test_pdv, jour="2026-10-17")
    code_origine = tour.code
    resp = await client.get("/api/exports/postier-planning",
                            params={"date": tour.date, "source": "ordonnancement"})
    wb = _classeur(resp.content)
    ws = wb["Tours"]
    entetes = _entetes(ws)
    ws.cell(_ligne(ws, entetes, tour.code), entetes["n° infolog"], "291189")
    await client.post(
        "/api/imports/tour-numbers",
        files={"file": ("ert.xlsx", _vers_fichier(wb),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    await db_session.refresh(tour)
    assert tour.code == code_origine
    assert tour.wms_tour_code == "291189"


@pytest.mark.asyncio
async def test_rejouer_l_import_ne_change_rien(client, db_session, test_region, test_pdv):
    """Un fichier partage se reimporte deux fois : la seconde doit etre neutre."""
    tour = await _tour(db_session, test_region, test_pdv, jour="2026-10-18")
    resp = await client.get("/api/exports/postier-planning",
                            params={"date": tour.date, "source": "ordonnancement"})
    wb = _classeur(resp.content)
    ws = wb["Tours"]
    entetes = _entetes(ws)
    ws.cell(_ligne(ws, entetes, tour.code), entetes["n° infolog"], "291190")
    fichier = _vers_fichier(wb)

    premier = await client.post("/api/imports/tour-numbers", files={"file": ("e.xlsx", fichier, "x")})
    second = await client.post("/api/imports/tour-numbers", files={"file": ("e.xlsx", fichier, "x")})
    assert premier.json()["updated"] == 1
    assert second.json()["updated"] == 0
    assert second.json()["unchanged"] == 1


@pytest.mark.asyncio
async def test_une_tournee_inconnue_est_signalee_pas_ignoree(
    client, db_session, test_region, test_pdv,
):
    """Si une ligne ne correspond a aucune tournee, l'exploitant doit le savoir
    — sinon il croit l'import complet alors qu'il manque des numeros."""
    tour = await _tour(db_session, test_region, test_pdv, jour="2026-10-19")
    resp = await client.get("/api/exports/postier-planning",
                            params={"date": tour.date, "source": "ordonnancement"})
    wb = _classeur(resp.content)
    ws = wb["Tours"]
    entetes = _entetes(ws)
    fantome = ws.max_row + 1
    ws.cell(fantome, entetes["code cmro"], "T-FANTOME")
    ws.cell(fantome, entetes["n° infolog"], "999999")
    ws.cell(_ligne(ws, entetes, tour.code), entetes["n° infolog"], "291191")

    resp = await client.post("/api/imports/tour-numbers",
                             files={"file": ("e.xlsx", _vers_fichier(wb), "x")})
    corps = resp.json()
    assert corps["updated"] == 1
    assert "T-FANTOME" in corps["not_found"]


@pytest.mark.asyncio
async def test_une_tournee_sans_numero_est_comptee(client, db_session, test_region, test_pdv):
    """Export reimporte sans rien saisir : rien ne doit etre ecrit, et le
    compteur doit dire pourquoi."""
    tour = await _tour(db_session, test_region, test_pdv, jour="2026-10-20")
    resp = await client.get("/api/exports/postier-planning",
                            params={"date": tour.date, "source": "ordonnancement"})
    resp = await client.post("/api/imports/tour-numbers",
                             files={"file": ("e.xlsx", resp.content, "x")})
    corps = resp.json()
    assert corps["updated"] == 0
    assert corps["without_number"] >= 1
    await db_session.refresh(tour)
    assert tour.wms_tour_code is None


@pytest.mark.asyncio
async def test_un_fichier_etranger_est_refuse(client):
    """Un classeur qui ne vient pas de CMRO n'a pas de code CMRO : le dire,
    plutot que de rapprocher au hasard."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = "n'importe quoi"
    resp = await client.post("/api/imports/tour-numbers",
                             files={"file": ("autre.xlsx", _vers_fichier(wb), "x")})
    assert resp.status_code == 400
    assert "export CMRO" in resp.json()["detail"]
